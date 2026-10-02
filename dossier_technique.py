"""
dossier_technique.py
=====================

Génère un « Dossier Technique » Excel (format standardisé, importé dans un logiciel
tiers) à partir d'un fichier de devis Excel.

IMPORTANT — SOURCE DES DONNÉES
------------------------------
La feuille « Bilan » du devis ne contient que des FORMULES (=IF(produit!B4=...)),
sans valeur calculée enregistrée dans le fichier (export Odoo / openpyxl).
Lue avec openpyxl (data_only=True), elle renvoie donc des cellules vides.
Le script lit désormais directement les feuilles SOURCES, qui contiennent de
vraies valeurs :
  - « produit »    : Client, Référence, Quantité (libellé en colonne A, valeur en B)
  - « Composants » : achats  (données à partir de la ligne 4)
  - « Operation »  : temps de fabrication (données à partir de la ligne 4)

CE QUI EST GÉNÉRÉ :
En-tête :
  - N° devis, N° dossier : laissés vides
  - Client <- produit (« Client »), Contact / Téléphone <- "Inconnu"
  - Référence <- produit (« Référence »), Quantité commandée <- produit (« Quantité »)

  - Descriptif (L26) <- produit!B8

Achats (feuille Composants) :
  - Catégorie dynamique "Achats <Section>" d'après la colonne A (SECTION)
  - Article=A, Support=C, Référence=D, Prix unitaire=J, Quantité=K, Coût devisé=L

Temps de fabrication (feuille Operation) :
  - Catégorie dynamique "Temps <Section>" d'après la colonne A (Section)
  - Article=A, Code=E, Nom opération=D, Tps=L

HYPOTHÈSES / LIMITES :
1. Le modèle cible est une mise en page fixe (sauts de page tous les 62 lignes).
   Les zones Achats et Temps sont effacées puis reconstruites dynamiquement.
2. Les champs hors cahier des charges sont laissés vides.
3. Si le contenu dépasse la capacité du modèle, le script s'arrête avec un
   message d'erreur explicite.
4. Le libellé d'unité (« €/<unité> ») est activable via INCLUDE_UNIT_LABEL.
"""
from __future__ import annotations

import copy
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# --------------------------------------------------------------------------
# Configuration / réglages ajustables
# --------------------------------------------------------------------------

INCLUDE_UNIT_LABEL = True  # affiche "€/<unité>" à côté du prix unitaire

# Feuilles du devis source (valeurs réelles, pas des formules)
PRODUIT_SHEET = "produit"
COMPOSANTS_SHEET = "Composants"
OPERATION_SHEET = "Operation"

# Cellule de « produit » recopiée telle quelle dans le Descriptif du dossier
PRODUIT_DESCRIPTIF_CELL = "B8"

# Libellés (colonne A de « produit ») -> valeur en colonne B
PRODUIT_LABEL_CLIENT = "Client"
PRODUIT_LABEL_REFERENCE = "Référence"
PRODUIT_LABEL_QUANTITE = "Quantité"

# Première ligne de données et nombre maximal de lignes lues
COMPOSANTS_FIRST_ROW = 4
OPERATION_FIRST_ROW = 4
MAX_LIGNES = 30

# Feuille « Composants » (achats)
COL_ACHAT_SECTION = 1     # A : SECTION
COL_ACHAT_SUPPORT = 3     # C : COMPOSANTS
COL_ACHAT_REF = 4         # D : CODE
COL_ACHAT_UNITE = 7       # G : UNITÉ
COL_ACHAT_PRIX_UNIT = 10  # J : P VENTE UNITAIRE
COL_ACHAT_QUANTITE = 11   # K : QUANTITÉ
COL_ACHAT_COUT_TOTAL = 12  # L : PRIX TOTAL

# Feuille « Operation » (temps)
COL_TEMPS_SECTION = 1     # A : Section
COL_TEMPS_OPERATION = 4   # D : Opération
COL_TEMPS_CODE = 5        # E : Code opération
COL_TEMPS_TPS = 12        # L : Temps total opérateur

# --------------------------------------------------------------------------
# Feuille cible ("DT Devis 1" dans le modèle fourni) : géométrie fixe
# --------------------------------------------------------------------------

TARGET_SHEET_NAME = "DT Devis 1"

# En-tête / informations générales
CELL_NDEVIS = "M8"
CELL_NDOSSIER = "M9"
CELL_TITRE_DOSSIER = "L5"       # "Dossier de fabrication - ..."
CELL_REF_TOP = "E2"             # petit rappel de référence en haut de page
CELL_CLIENT = "M11"
CELL_CONTACT = "M12"
CELL_TELEPHONE = "M13"
CELL_REFERENCE = "M14"
CELL_QUANTITE_CMD = "N16"
CELL_DESCRIPTIF = "L26"          # Descriptif <- produit!B8

# Doublons littéraux du N° devis affichés sur chaque page (pas des formules)
NDEVIS_ECHO_CELLS = ["M70", "M132", "M194", "M256", "M318"]

# Champs hors périmètre du mapping -> vidés pour ne pas polluer le dossier
CELLS_TO_BLANK = [
    "D8", "D9", "D10",          # Départ / Franco / Catégorie
    "O13", "P13",                # Email
    "L20", "P20", "L21", "L22",  # Fichiers / BAT
    "L56",                       # Observations
    "L61", "N61",                # Livraison prévue
    "L63",                       # Adresse de livraison
    "P324", "P325", "P326", "P327", "P329", "Q329",
    "P330", "P331", "P332", "P333",
    "P339", "P340", "P341", "P342", "P343", "P344", "Q344",
    "P345", "P346", "P347", "P348",
    "Q354", "Q355", "Q356", "Q357", "Q358", "Q359",
]

# Colonnes utilisées pour les blocs "Achats" / "Temps" (L à T)
BLOCK_COL_FIRST = 12  # L
BLOCK_COL_LAST = 20   # T (la fusion d'en-tête va jusqu'à S=19 dans le modèle)
HEADER_MERGE_LAST_COL = 19  # S

# --- Zone ACHATS : pagination fixe du modèle --------------------------------
ACHATS_PAGE1_START = 73
ACHATS_DEAD_ZONE = (124, 135)   # (exclu, cible) : > 124 et < 135 -> saute à 135
ACHATS_PAGE2_START = 135
ACHATS_HARD_MAX = 190           # dernière ligne utilisable avant "Temps de fabrication"
ACHATS_ITEM_BLOCK_ROWS = 3
ACHATS_ITEM_STEP = 4            # 3 lignes de données + 1 ligne vide
ACHATS_CLEAR_RANGES = [(73, 128), (135, 190)]
# lignes "modèles" dont on copie le style
ACHATS_HEADER_STYLE_ROW = 73
ACHATS_ITEM_STYLE_ROWS = (94, 95, 96)  # bloc "Achats des conditionnements"

# --- Zone TEMPS : pagination fixe du modèle ---------------------------------
TEMPS_PAGE1_START = 197
TEMPS_DEAD_ZONE = (244, 259)
TEMPS_PAGE2_START = 259
TEMPS_HARD_MAX = 313            # dernière ligne utilisable avant "Synthèse"
TEMPS_ITEM_STEP = 1
TEMPS_CLEAR_RANGES = [(197, 251), (259, 313)]
TEMPS_HEADER_STYLE_ROW = 197
TEMPS_ITEM_STYLE_ROW = 198


# --------------------------------------------------------------------------
# Réparation des fichiers .xlsx/.xlsm mal formés (data validation invalide)
# --------------------------------------------------------------------------

def repair_xlsx(src: str | Path, dst: str | Path) -> Path:
    """Corrige les attributs errorStyle="" invalides que certains exports
    Excel/no-code laissent dans les data validations, et qui font planter
    openpyxl ("Value must be one of {'stop','information','warning'}")."""
    src, dst = Path(src), Path(dst)
    shutil.copyfile(src, dst)
    tmp = dst.with_suffix(dst.suffix + ".tmp")
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                text = data.decode("utf-8", errors="replace")
                text = re.sub(
                    r'errorStyle="([^"]*)"',
                    lambda m: 'errorStyle="stop"' if m.group(1) not in ("stop", "information", "warning") else m.group(0),
                    text,
                )
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(dst)
    return dst


def safe_load_workbook(path: str | Path, data_only: bool = True):
    """Charge un classeur avec openpyxl ; si la lecture échoue à cause d'une
    data validation corrompue, répare une copie temporaire puis réessaie."""
    path = Path(path)
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except ValueError as exc:
        if "could not read worksheets" not in str(exc):
            raise
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            repaired = Path(tmp_dir) / f"repaired_{path.name}"
            repair_xlsx(path, repaired)
            return openpyxl.load_workbook(repaired, data_only=data_only)


# --------------------------------------------------------------------------
# Lecture du devis
# --------------------------------------------------------------------------

@dataclass
class AchatLigne:
    article: str
    support: str
    reference: str
    unite: str
    prix_unitaire: Optional[float]
    quantite: Optional[float]
    cout_devise: Optional[float]


@dataclass
class TempsLigne:
    article: str
    code: str
    nom_operation: str
    tps: Optional[float]


@dataclass
class DevisData:
    client: str = ""
    reference: str = ""
    quantite: Optional[float] = None
    descriptif: str = ""
    achats: dict[str, list[AchatLigne]] = field(default_factory=dict)   # section -> lignes, ordre d'apparition
    temps: dict[str, list[TempsLigne]] = field(default_factory=dict)


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def _produit_value(ws, label: str):
    """Retourne la valeur (colonne B) de la ligne dont le libellé (colonne A) vaut `label`."""
    wanted = label.strip().lower()
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if isinstance(row[0], str) and row[0].strip().lower() == wanted:
            return row[1]
    return None


def read_devis(path: str | Path) -> DevisData:
    """Lit le devis (feuilles produit / Composants / Operation) et retourne les
    données structurées nécessaires à la génération du Dossier Technique."""
    wb = safe_load_workbook(path, data_only=True)
    for name in (PRODUIT_SHEET, COMPOSANTS_SHEET, OPERATION_SHEET):
        if name not in wb.sheetnames:
            raise ValueError(
                f"La feuille « {name} » est introuvable dans {path}. "
                f"Feuilles disponibles : {wb.sheetnames}"
            )
    wp, wc, wo = wb[PRODUIT_SHEET], wb[COMPOSANTS_SHEET], wb[OPERATION_SHEET]

    data = DevisData(
        client=_clean(_produit_value(wp, PRODUIT_LABEL_CLIENT)),
        reference=_clean(_produit_value(wp, PRODUIT_LABEL_REFERENCE)),
        quantite=_produit_value(wp, PRODUIT_LABEL_QUANTITE),
        descriptif=wp[PRODUIT_DESCRIPTIF_CELL].value or "",
    )

    # --- Achats (feuille Composants) ---
    for row in range(COMPOSANTS_FIRST_ROW, COMPOSANTS_FIRST_ROW + MAX_LIGNES):
        if not _clean(wc.cell(row=row, column=COL_ACHAT_SUPPORT).value):
            continue  # pas de composant -> ligne ignorée
        section = _clean(wc.cell(row=row, column=COL_ACHAT_SECTION).value)
        ligne = AchatLigne(
            article=section,
            support=_clean(wc.cell(row=row, column=COL_ACHAT_SUPPORT).value),
            reference=_clean(wc.cell(row=row, column=COL_ACHAT_REF).value),
            unite=_clean(wc.cell(row=row, column=COL_ACHAT_UNITE).value),
            prix_unitaire=wc.cell(row=row, column=COL_ACHAT_PRIX_UNIT).value,
            quantite=wc.cell(row=row, column=COL_ACHAT_QUANTITE).value,
            cout_devise=wc.cell(row=row, column=COL_ACHAT_COUT_TOTAL).value,
        )
        data.achats.setdefault(section, []).append(ligne)

    # --- Temps (feuille Operation) ---
    for row in range(OPERATION_FIRST_ROW, OPERATION_FIRST_ROW + MAX_LIGNES):
        if not _clean(wo.cell(row=row, column=COL_TEMPS_OPERATION).value):
            continue
        section = _clean(wo.cell(row=row, column=COL_TEMPS_SECTION).value)
        ligne = TempsLigne(
            article=section,
            code=_clean(wo.cell(row=row, column=COL_TEMPS_CODE).value),
            nom_operation=_clean(wo.cell(row=row, column=COL_TEMPS_OPERATION).value),
            tps=wo.cell(row=row, column=COL_TEMPS_TPS).value,
        )
        data.temps.setdefault(section, []).append(ligne)

    return data


# --------------------------------------------------------------------------
# Écriture du Dossier Technique
# --------------------------------------------------------------------------

def _copy_style(src_cell, dst_cell):
    dst_cell.font = copy.copy(src_cell.font)
    dst_cell.border = copy.copy(src_cell.border)
    dst_cell.fill = copy.copy(src_cell.fill)
    dst_cell.alignment = copy.copy(src_cell.alignment)
    dst_cell.number_format = src_cell.number_format
    dst_cell.protection = copy.copy(src_cell.protection)


def _clear_zone(ws: Worksheet, ranges: list[tuple[int, int]]):
    """Vide les valeurs et défusionne les cellules dans les plages de lignes
    données (bornes incluses), sur toutes les colonnes utilisées."""
    max_col = ws.max_column
    for r1, r2 in ranges:
        for mr in list(ws.merged_cells.ranges):
            if r1 <= mr.min_row and mr.max_row <= r2:
                ws.unmerge_cells(str(mr))
        for row in range(r1, r2 + 1):
            for col in range(1, max_col + 1):
                ws.cell(row=row, column=col).value = None


def _gate(row: int, dead_zone: tuple[int, int], jump_to: int, hard_max: int, label: str) -> int:
    lo, hi = dead_zone
    if lo < row < hi:
        row = jump_to
    if row > hard_max:
        raise OverflowError(
            f"Le contenu du dossier technique ({label}) dépasse la capacité du "
            f"modèle (ligne {row} > {hard_max}). Réduisez le nombre de lignes "
            f"« {label} » dans le devis, répartissez-les en moins de catégories, "
            f"ou étendez le modèle Excel avec des pages supplémentaires."
        )
    return row


def _write_achat_header(ws: Worksheet, row: int, title: str, style_row: int):
    ws.merge_cells(start_row=row, start_column=BLOCK_COL_FIRST, end_row=row, end_column=HEADER_MERGE_LAST_COL)
    cell = ws.cell(row=row, column=BLOCK_COL_FIRST)
    src = ws.cell(row=style_row, column=BLOCK_COL_FIRST)
    _copy_style(src, cell)
    cell.value = title
    if style_row in ws.row_dimensions and ws.row_dimensions[style_row].height:
        ws.row_dimensions[row].height = ws.row_dimensions[style_row].height


def _write_achat_item(ws: Worksheet, row: int, item: AchatLigne):
    style_rows = ACHATS_ITEM_STYLE_ROWS  # (label_row, support_row, prix_row) offsets 0,1,2

    def put(r_off, col, value, style_col=None):
        style_col = style_col or col
        c = ws.cell(row=row + r_off, column=col)
        _copy_style(ws.cell(row=style_rows[r_off], column=style_col), c)
        c.value = value

    # ligne 1 : Article / Référence
    put(0, 12, "Article : ")
    put(0, 13, item.article)
    put(0, 18, "Référence : ")
    put(0, 19, item.reference)

    # ligne 2 : Support / Surface (vide) / Fournisseur (vide)
    put(1, 12, "Support : ")
    put(1, 13, item.support)
    put(1, 15, "Surface : ")
    put(1, 16, None)
    put(1, 18, "Fournisseur : ")
    put(1, 19, None)

    # ligne 3 : Prix unitaire / Quantité / Coût devisé
    put(2, 12, "Prix unitaire : ")
    put(2, 13, item.prix_unitaire)
    if INCLUDE_UNIT_LABEL and item.unite:
        put(2, 14, f"€/{item.unite}")
    else:
        put(2, 14, None)
    put(2, 15, "Quantité : ")
    put(2, 16, item.quantite)
    put(2, 18, "Coût devisé : ")
    put(2, 19, item.cout_devise)


def _write_temps_header(ws: Worksheet, row: int, title: str, style_row: int):
    ws.merge_cells(start_row=row, start_column=BLOCK_COL_FIRST, end_row=row, end_column=HEADER_MERGE_LAST_COL)
    cell = ws.cell(row=row, column=BLOCK_COL_FIRST)
    src = ws.cell(row=style_row, column=BLOCK_COL_FIRST)
    _copy_style(src, cell)
    cell.value = title
    if style_row in ws.row_dimensions and ws.row_dimensions[style_row].height:
        ws.row_dimensions[row].height = ws.row_dimensions[style_row].height


def _write_temps_item(ws: Worksheet, row: int, item: TempsLigne):
    style_row = TEMPS_ITEM_STYLE_ROW

    def put(col, value):
        c = ws.cell(row=row, column=col)
        _copy_style(ws.cell(row=style_row, column=col), c)
        c.value = value

    put(12, f"Article : {item.article}")
    put(14, f"Code : {item.code}")
    put(15, item.nom_operation)
    put(19, item.tps)


def _write_achats(ws: Worksheet, achats: dict[str, list[AchatLigne]]):
    row = ACHATS_PAGE1_START
    for section, items in achats.items():
        title = f"Achats {section}"
        _write_achat_header(ws, row, title, ACHATS_HEADER_STYLE_ROW)
        item_row = row + 1
        for item in items:
            item_row = _gate(item_row, ACHATS_DEAD_ZONE, ACHATS_PAGE2_START, ACHATS_HARD_MAX, "achats")
            _write_achat_item(ws, item_row, item)
            item_row += ACHATS_ITEM_STEP
        n = len(items)
        next_header = row + 4 * n + 2 if n else row + 2
        row = _gate(next_header, ACHATS_DEAD_ZONE, ACHATS_PAGE2_START, ACHATS_HARD_MAX, "achats")


def _write_temps(ws: Worksheet, temps: dict[str, list[TempsLigne]]):
    row = TEMPS_PAGE1_START
    for section, items in temps.items():
        title = f"Temps {section}"
        _write_temps_header(ws, row, title, TEMPS_HEADER_STYLE_ROW)
        item_row = row + 1
        for item in items:
            item_row = _gate(item_row, TEMPS_DEAD_ZONE, TEMPS_PAGE2_START, TEMPS_HARD_MAX, "temps de fabrication")
            _write_temps_item(ws, item_row, item)
            item_row += TEMPS_ITEM_STEP
        n = len(items)
        next_header = row + n + 2
        row = _gate(next_header, TEMPS_DEAD_ZONE, TEMPS_PAGE2_START, TEMPS_HARD_MAX, "temps de fabrication")


def generate_dossier_technique(
    devis_path: str | Path,
    template_path: str | Path,
    output_path: str | Path,
) -> Path:
    """Génère le Dossier Technique à partir d'un devis et du modèle Excel,
    et l'enregistre dans output_path."""
    devis = read_devis(devis_path)

    wb = safe_load_workbook(template_path, data_only=False)
    if TARGET_SHEET_NAME not in wb.sheetnames:
        # on prend la première feuille si le nom diffère (modèle personnalisé)
        ws = wb[wb.sheetnames[0]]
    else:
        ws = wb[TARGET_SHEET_NAME]

    # 1) informations générales
    ws[CELL_NDEVIS] = None
    ws[CELL_NDOSSIER] = None
    for coord in NDEVIS_ECHO_CELLS:
        ws[coord] = None
    ws[CELL_TITRE_DOSSIER] = "Dossier de fabrication"
    ws[CELL_REF_TOP] = devis.reference
    ws[CELL_CLIENT] = devis.client
    ws[CELL_CONTACT] = "Inconnu"
    ws[CELL_TELEPHONE] = "Inconnu"
    ws[CELL_REFERENCE] = devis.reference
    ws[CELL_QUANTITE_CMD] = devis.quantite

    # 2) champs hors périmètre -> vidés
    for coord in CELLS_TO_BLANK:
        ws[coord] = None

    # 2 bis) Descriptif <- produit!B8 (valeur brute, retours à la ligne conservés)
    ws[CELL_DESCRIPTIF] = devis.descriptif

    # 3) zones achats / temps : on efface puis on reconstruit dynamiquement
    _clear_zone(ws, ACHATS_CLEAR_RANGES)
    _clear_zone(ws, TEMPS_CLEAR_RANGES)
    _write_achats(ws, devis.achats)
    _write_temps(ws, devis.temps)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


# --------------------------------------------------------------------------
# Ligne de commande
# --------------------------------------------------------------------------

def _default_template() -> Path:
    return Path(__file__).parent / "template_dossier_technique.xlsx"


def main(argv=None):
    import argparse
    import sys

    parser = argparse.ArgumentParser(
        description="Génère un Dossier Technique Excel à partir d'un nouveau devis."
    )
    parser.add_argument("devis", help="Chemin du fichier de devis (.xlsx) à traiter")
    parser.add_argument(
        "-o", "--output", default=None,
        help="Chemin du fichier Dossier Technique à générer "
             "(par défaut : <nom_du_devis>_DT.xlsx à côté du devis)",
    )
    parser.add_argument(
        "-t", "--template", default=None,
        help="Chemin du modèle Excel cible à utiliser "
             "(par défaut : template_dossier_technique.xlsx fourni avec le script)",
    )
    args = parser.parse_args(argv)

    devis_path = Path(args.devis)
    template_path = Path(args.template) if args.template else _default_template()
    output_path = Path(args.output) if args.output else devis_path.with_name(devis_path.stem + "_DT.xlsx")

    if not template_path.exists():
        print(f"Modèle introuvable : {template_path}", file=sys.stderr)
        sys.exit(1)

    try:
        out = generate_dossier_technique(devis_path, template_path, output_path)
    except OverflowError as exc:
        print(f"Erreur : {exc}", file=sys.stderr)
        sys.exit(2)
    except Exception as exc:  # noqa: BLE001
        print(f"Erreur lors de la génération : {exc}", file=sys.stderr)
        sys.exit(1)

    print(f"Dossier Technique généré : {out}")


if __name__ == "__main__":
    main()